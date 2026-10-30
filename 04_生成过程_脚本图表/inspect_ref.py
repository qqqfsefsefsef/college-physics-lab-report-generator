import sys, json
from docx import Document
from docx.shared import Pt

for path in [r'C:\Users\lidey\Desktop\实验报告制作沙盒\热膨胀系数.docx',
             r'C:\Users\lidey\Desktop\实验报告制作沙盒\霍尔效应-已填数据.docx']:
    print('=' * 90)
    print('FILE:', path)
    d = Document(path)
    sec = d.sections[0]
    print('page w/h (cm):', round(sec.page_width.cm, 2), round(sec.page_height.cm, 2))
    print('margins L/R/T/B (cm):', round(sec.left_margin.cm, 2), round(sec.right_margin.cm, 2),
          round(sec.top_margin.cm, 2), round(sec.bottom_margin.cm, 2))
    print('n_paragraphs:', len(d.paragraphs), 'n_tables:', len(d.tables))
    print('--- styles used:', sorted({p.style.name for p in d.paragraphs}))
    print('--- PARAGRAPHS ---')
    for i, p in enumerate(d.paragraphs):
        t = p.text.strip()
        runs = p.runs
        f = ''
        if runs:
            r0 = runs[0]
            f = f"{r0.font.name}|{r0.font.size.pt if r0.font.size else None}|b={r0.bold}|i={r0.italic}"
            ea = r0._element.rPr.rFonts.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia') if r0._element.rPr is not None and r0._element.rPr.rFonts is not None else None
            f += f"|ea={ea}"
        print(f'[{i}] <{p.style.name}> align={p.alignment} ind_first={p.paragraph_format.first_line_indent} | {f} | {t[:120]}')
    print('--- TABLES ---')
    for ti, tb in enumerate(d.tables):
        print(f'TABLE {ti}: style={tb.style.name if tb.style else None} rows={len(tb.rows)} cols={len(tb.columns)} align={tb.alignment}')
        for ri, row in enumerate(tb.rows[:4]):
            cells = [c.text.strip().replace('\n', ' / ')[:28] for c in row.cells]
            print('   r%d: %s' % (ri, ' || '.join(cells)))
    print('--- INLINE SHAPES:', len(d.inline_shapes))
    for sh in d.inline_shapes:
        print('   ', sh.width.cm, 'x', sh.height.cm, sh.type)
