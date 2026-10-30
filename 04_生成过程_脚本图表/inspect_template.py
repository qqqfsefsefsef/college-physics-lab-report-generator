# -*- coding: utf-8 -*-
import io
from docx import Document
from docx.oxml.ns import qn
from lxml import etree

buf = io.StringIO()
def p(x): buf.write(str(x) + '\n')

# header / footer XML of reference
for path in [r'C:\Users\lidey\Desktop\实验报告制作沙盒\热膨胀系数.docx']:
    d = Document(path)
    s = d.sections[0]
    for name, part in [('HEADER', s.header), ('FOOTER', s.footer)]:
        p(f'--- {name} part type={part.is_linked_to_previous}')
        for i, para in enumerate(part.paragraphs):
            xml = etree.tostring(para._p, encoding='unicode')
            p(f'  p{i}: {xml}')
        for t in part.tables:
            p(f'  table rows={len(t.rows)}')
    p('--- sectPr: ' + etree.tostring(s._sectPr, encoding='unicode')[:2000])

# template
tp = r'C:\Users\lidey\Desktop\实验报告制作沙盒\附录2 大学物理实验电子报告模版及范例.docx'
d = Document(tp)
s = d.sections[0]
p('=' * 90)
p('TEMPLATE 附录2')
p(f'page {s.page_width.cm:.2f}x{s.page_height.cm:.2f} margins L{s.left_margin.cm:.2f} R{s.right_margin.cm:.2f} T{s.top_margin.cm:.2f} B{s.bottom_margin.cm:.2f}')
p(f'header: {[x.text for x in s.header.paragraphs]}  footer: {[x.text for x in s.footer.paragraphs]}')
p('styles: ' + str([st.name for st in d.styles if st.type is not None][:40]))
from docx.table import Table
from docx.text.paragraph import Paragraph
i = 0
for child in d.element.body.iterchildren():
    if child.tag == qn('w:p'):
        par = Paragraph(child, d)
        nimg = len(child.findall('.//' + qn('w:drawing')))
        # detect colored runs
        cols = set()
        for r in par.runs:
            c = r.font.color
            if c is not None and c.rgb is not None:
                cols.add(str(c.rgb))
        p(f'P{i} <{par.style.name}> ind1st={par.paragraph_format.first_line_indent} nimg={nimg} colors={cols} :: {par.text}')
        i += 1
    elif child.tag == qn('w:tbl'):
        t = Table(child, d)
        p(f'TBL rows={len(t.rows)} cols={len(t.columns)}: ' + ' | '.join(c.text[:20] for c in t.rows[0].cells)[:200])
p('n_paragraphs=%s n_tables=%s n_images=%s' % (len(d.paragraphs), len(d.tables), len(d.inline_shapes)))

open(r'C:\Users\lidey\Desktop\实验报告制作沙盒\template_dump.txt', 'w', encoding='utf-8').write(buf.getvalue())
print('ok', len(buf.getvalue()))

