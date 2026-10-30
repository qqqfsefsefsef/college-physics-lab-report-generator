# -*- coding: utf-8 -*-
"""Build 钨的逸出功测定.docx from 热膨胀系数.docx skeleton + content_tungsten.md + data_tungsten.json."""
import json
import os
import re

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

# 路径可由流水线用环境变量覆盖（默认保持原有行为）
W = os.environ.get('DSH_WORKSPACE', r'C:\Users\lidey\Desktop\实验报告制作沙盒')
GEN = os.environ.get('PIPE_GEN', W + r'\04_生成过程_脚本图表')
BASE = os.environ.get('PIPE_BASE', W + r'\03_参考模板_旧报告\热膨胀系数.docx')
OUT = os.environ.get('PIPE_OUT', W + r'\05_成品_报告\钨的逸出功测定.docx')
SHOT = os.environ.get('PIPE_SHOT', W + r'\01_素材_数据与截图\Screenshot_2026-10-26-22-19-51-481_com.chaoxing.m.jpg')
MD = os.environ.get('PIPE_MD', GEN + r'\content_tungsten.md')
DATA = os.environ.get('PIPE_DATA', GEN + r'\data_tungsten.json')

GRAY = '494949'
RED = 'FF0000'
CN = '宋体'
EN = 'Times New Roman'

D = json.load(open(DATA, encoding='utf-8'))
T1 = D['table1']
T2 = D['table2']
T3 = D['table3']
F3FIT = D['table3_fit']
LG_IS = D['lg_is_from_extrapolation']
CONST = D['meta']['constants']

U_A = T1['u_a_V']
I_F = T1['i_f_A']
I_A = T1['i_a_uA']
SQRT_U = T2['sqrt_u_a']
LG_IA = T2['lg_i_a']
ROW_T216 = T2['lg_i_a_row_T216_extra']
T_LIST = T3['T_1000K']
INV_T = T3['inv_T_1e3_per_K']
LG_I = T3['lg_i']
LG_IT2 = T3['lg_is_over_T2_written']

K_B = CONST['k_J_per_K']
E_C = CONST['e_C']
PHI_ACCEPT = CONST['phi_accepted_eV']

M_PLOT = F3FIT['slope']
B_PLOT = F3FIT['intercept']
R_VAL = F3FIT['r']
R2_VAL = F3FIT['r2']
SE_M = F3FIT['se_slope']

LN10 = 2.303                                # 与讲义式（3）的写法一致
M_TRUE = M_PLOT * 1e3                       # K
PHI_J = LN10 * K_B * abs(M_TRUE)            # J
PHI_EV = PHI_J / E_C                        # eV
U_PHI_EV = LN10 * K_B * SE_M * 1e3 / E_C
REL_ERR = (PHI_EV - PHI_ACCEPT) / PHI_ACCEPT * 100.0

# ---------------------------------------------------------------- content
SECTIONS = {}
_SRC = open(MD, encoding='utf-8').read()
for _blk in re.split(r'^===\s*(.+?)\s*===\s*$', _SRC, flags=re.M)[1:]:
    pass
_parts = re.split(r'^===\s*(.+?)\s*===\s*$', _SRC, flags=re.M)
for _i in range(1, len(_parts), 2):
    SECTIONS[_parts[_i].strip()] = [ln.strip() for ln in _parts[_i + 1].splitlines() if ln.strip()]

FORMULAS = {
    'F1': '*f*(*E*) = d*N*/d*E* = (4π/*h*³)(2*m*)^{3/2}·*E*^{1/2}·[exp((*E* − *E*_{F})/*kT*) + 1]^{−1}',
    'F2': '*E*_{0} = *E*_{b} − *E*_{F} = *eφ*',
    'F3': '*I* = *AST*²·exp(−*eφ*/*kT*)',
    'F4': 'lg(*I*/*T*²) = lg *AS* − *eφ*/(2.30*kT*) = lg *AS* − 5.04×10³·*φ*/*T*',
    'F5': '*I*_{a} = *I*·exp(0.439√*E*_{a}/*T*)',
    'F6': 'lg *I*_{a} = lg *I* + 0.439√*E*_{a}/(2.30*T*)',
    'F7': '*E*_{a} = *U*_{a}/(*r*_{1}·ln(*r*_{2}/*r*_{1}))',
    'F8': 'lg *I*_{a} = lg *I* + [0.439/(2.30*T*)]·√(*U*_{a}/(*r*_{1}ln(*r*_{2}/*r*_{1})))',
    'F9': '*T* = 920.0 + 1600·*I*_{f}',
}

doc = Document(BASE)


# ---------------------------------------------------------------- red notes
def red_text(idx):
    par = doc.paragraphs[idx]
    out = []
    for r in par.runs:
        col = r.font.color
        if col is not None and col.rgb is not None and str(col.rgb) == RED:
            out.append(r.text)
    return ''.join(out)


NOTE_PRINCIPLE = red_text(35)
NOTE_STEPS = red_text(72)
NOTE_DATA = red_text(87)
NOTE_EXP = red_text(91)

# ---------------------------------------------------------------- truncate
pars = doc.paragraphs
assert pars[27].text.strip().startswith('实验名称'), pars[27].text
for r in pars[27].runs:
    if r.text.strip() == '热膨胀系数':
        r.text = '钨的逸出功测定'

keep = pars[29]._p
body_el = doc.element.body
seen = False
for child in list(body_el):
    if child is keep:
        seen = True
        continue
    if not seen:
        continue
    if child.tag == qn('w:sectPr'):
        continue
    body_el.remove(child)

# ---------------------------------------------------------------- primitives
ALIGN = {'left': WD_ALIGN_PARAGRAPH.LEFT, 'center': WD_ALIGN_PARAGRAPH.CENTER,
         'right': WD_ALIGN_PARAGRAPH.RIGHT, 'justify': WD_ALIGN_PARAGRAPH.JUSTIFY}


def clean(s):
    s = re.sub(r'\s+([，。、；：？！）】》”])', r'\1', s)
    s = re.sub(r'([（【“])\s+', r'\1', s)
    s = re.sub(r'\s{2,}', ' ', s)
    return s.strip()


def mk_run(p, text, size=12, bold=False, color=GRAY, italic=False,
           sup=False, sub=False):
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = RGBColor.from_string(color)
    rPr = r._r.get_or_add_rPr()
    rf = rPr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts')
        rPr.insert(0, rf)
    rf.set(qn('w:ascii'), EN)
    rf.set(qn('w:hAnsi'), EN)
    rf.set(qn('w:eastAsia'), CN)
    rf.set(qn('w:cs'), EN)
    if sup:
        r.font.superscript = True
    if sub:
        r.font.subscript = True
    return r


def new_p(align='left', indent_chars=0, space_before=0, space_after=0,
          keep_next=False):
    p = doc.add_paragraph()
    p.alignment = ALIGN[align]
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.widow_control = True
    if indent_chars:
        pf.first_line_indent = Pt(12 * indent_chars)
    if keep_next:
        pf.keep_with_next = True
    return p


def body(text, indent=True, size=12, bold=False, color=GRAY, align='left',
         keep_next=False):
    p = new_p(align=align, indent_chars=2 if indent else 0, keep_next=keep_next)
    mk_run(p, clean(text), size=size, bold=bold, color=color)
    return p


def heading(name, pct='', note=None):
    p = new_p(keep_next=True)
    mk_run(p, '【', bold=False)
    mk_run(p, name, bold=True)
    mk_run(p, '】' + pct, bold=False)
    if note:
        mk_run(p, note, color=RED)
    return p


_TOK = re.compile(r'\*([^*]+)\*|_\{([^{}]*)\}|\^\{([^{}]*)\}')


def formula(expr):
    p = new_p(align='center', space_before=6, space_after=6)
    pos = 0
    for m in _TOK.finditer(expr):
        if m.start() > pos:
            mk_run(p, expr[pos:m.start()])
        if m.group(1) is not None:
            mk_run(p, m.group(1), italic=True)
        elif m.group(2) is not None:
            mk_run(p, m.group(2), sub=True)
        else:
            mk_run(p, m.group(3), sup=True)
        pos = m.end()
    if pos < len(expr):
        mk_run(p, expr[pos:])
    return p


def tw(n):
    return Cm(n / 1440.0 * 2.54)


def cell_setup(cell, width_tw, size=10.5, bold=False, align='center', borders=True):
    cell.width = tw(width_tw)
    if borders:
        tcPr = cell._tc.get_or_add_tcPr()
        b = OxmlElement('w:tcBorders')
        for edge in ('top', 'left', 'bottom', 'right'):
            e = OxmlElement('w:' + edge)
            e.set(qn('w:val'), 'single')
            e.set(qn('w:sz'), '4')
            e.set(qn('w:space'), '0')
            e.set(qn('w:color'), 'auto')
            b.append(e)
        tcPr.append(b)
        mar = OxmlElement('w:tcMar')
        for edge, v in (('top', 0), ('left', 30), ('bottom', 0), ('right', 30)):
            e = OxmlElement('w:' + edge)
            e.set(qn('w:w'), str(v))
            e.set(qn('w:type'), 'dxa')
            mar.append(e)
        tcPr.append(mar)


def put(cell, text, size=10.5, bold=False, align='center'):
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = cell.paragraphs[0]
    p.alignment = ALIGN[align]
    pf = p.paragraph_format
    pf.first_line_indent = Pt(0)
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.widow_control = True
    mk_run(p, text, size=size, bold=bold)


def put_lines(cell, lines, size=10.5, bold=False, align='center'):
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    first = cell.paragraphs[0]
    for r in list(first.runs):
        r._r.getparent().remove(r._r)
    for i, line in enumerate(lines):
        p = first if i == 0 else cell.add_paragraph()
        p.alignment = ALIGN[align]
        pf = p.paragraph_format
        pf.first_line_indent = Pt(0)
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.widow_control = True
        mk_run(p, line, size=size, bold=bold)


def new_table(rows, cols, widths):
    tbl = doc.add_table(rows=rows, cols=cols)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    for i, wv in enumerate(widths):
        tbl.columns[i].width = tw(wv)
    for i, wv in enumerate(widths):
        for r in tbl.rows:
            r.cells[i].width = tw(wv)
    return tbl


def row_height(tbl, idx, val_tw, rule='atLeast'):
    tr = tbl.rows[idx]._tr
    trPr = tr.get_or_add_trPr()
    h = OxmlElement('w:trHeight')
    h.set(qn('w:val'), str(val_tw))
    h.set(qn('w:hRule'), rule)
    trPr.append(h)


def kill_vmerge_inner_lines(tbl):
    for tc in tbl._tbl.iter(qn('w:tc')):
        tcPr = tc.tcPr
        if tcPr is None:
            continue
        v = tcPr.find(qn('w:vMerge'))
        if v is None or v.get(qn('w:val')) not in (None, 'continue'):
            continue
        b = tcPr.find(qn('w:tcBorders'))
        if b is None:
            continue
        top = b.find(qn('w:top'))
        if top is not None:
            top.set(qn('w:val'), 'nil')


def figure(path, width_cm, caption=None):
    p = new_p(align='center', space_before=6, space_after=3, keep_next=True)
    p.add_run().add_picture(path, width=Cm(width_cm))
    if caption:
        cp = new_p(align='center', space_after=6)
        mk_run(cp, caption, size=10.5, bold=True)


def caption(text):
    p = new_p(align='center', space_before=6, space_after=3, keep_next=True)
    mk_run(p, text, size=10.5, bold=True)
    return p


def para_seq(lines):
    """Render a list of md lines as paragraphs / formulas."""
    for ln in lines:
        m = re.fullmatch(r'\[\[(F\d+)\]\]', ln)
        if m:
            formula(FORMULAS[m.group(1)])
        else:
            body(ln)


# ---------------------------------------------------------------- 实验目的
heading('实验目的', '10%')
body('1、了解有关热电子发射的基本规律，认识热电子发射与逸出功之间的联系。')
body('2、用里查孙直线法测定钨丝的电子逸出电势 φ，并求出逸出功 eφ。')
body('3、进一步学习数据处理方法，掌握把非线性关系转化为线性关系、用外推法求取'
     '不易直接测量的物理量的思路。')

# ---------------------------------------------------------------- 实验仪器
heading('实验仪器')
_inst = SECTIONS['仪器']
_fl = [i for i, l in enumerate(_inst) if re.fullmatch(r'\[\[(F\d+)\]\]', l)]
_cut = (_fl[0] if _fl else len(_inst))
para_seq(_inst[:_cut])
figure(GEN + r'\fig_w_wiring.png', 12.60, '图 1  实验线路原理示意图')
para_seq(_inst[_cut:])
figure(GEN + r'\fig_w_tcal.png', 12.60, '图 2  阴极温度与灯丝电流的关系')
caption('附表 1  阴极灯丝温度 T 与灯丝电流 I_f 的关系（厂家定标）')
_tw_a = [1700] + [943] * 7
_tcal = new_table(2, 8, _tw_a)
for i, (lab, vals) in enumerate([
        ('I_f / A', ['0.50', '0.55', '0.60', '0.65', '0.70', '0.75', '0.80']),
        ('T / 10³K', ['1.72', '1.80', '1.88', '1.96', '2.04', '2.12', '2.20'])]):
    cell_setup(_tcal.cell(i, 0), _tw_a[0])
    put(_tcal.cell(i, 0), lab, bold=True)
    for j, v in enumerate(vals):
        cell_setup(_tcal.cell(i, j + 1), _tw_a[j + 1])
        put(_tcal.cell(i, j + 1), v)

# ---------------------------------------------------------------- 实验原理
heading('实验原理', '10%', NOTE_PRINCIPLE)
for _sub, _sec in (('（一）金属电子的逸出功', '原理1'),
                   ('（二）热电子发射公式', '原理2'),
                   ('（三）里查孙直线法', '原理3'),
                   ('（四）加速电场外延法与肖脱基效应', '原理4')):
    body(_sub, indent=False, bold=True, keep_next=True)
    para_seq(SECTIONS[_sec])

# ---------------------------------------------------------------- 步骤
heading('实验内容与步骤', '10%', NOTE_STEPS)
para_seq(SECTIONS['步骤'])

# ---------------------------------------------------------------- 原始数据照片
heading('有教师签字的原始数据记录照片（请竖排粘贴）', '10%')
figure(GEN + r'\data_photo_w.jpg', 12.40)

# ---------------------------------------------------------------- 数据处理
heading('数据处理及结果', '20%', NOTE_DATA)
para_seq(SECTIONS['数据说明'])

caption('表 1  不同灯丝电流 I_f 和阳极电压 U_a 对应的阳极电流 I_a 值')
_w1 = [1180] + [890] * 8
_tb1 = new_table(9, 9, _w1)
cell_setup(_tb1.cell(0, 0), _w1[0])
put_lines(_tb1.cell(0, 0), ['U_a / V', 'I_a / μA', 'I_f / A'], size=9)
for j, ua in enumerate(U_A):
    cell_setup(_tb1.cell(0, j + 1), _w1[j + 1])
    put(_tb1.cell(0, j + 1), str(ua), bold=True)
for i, ifv in enumerate(I_F):
    cell_setup(_tb1.cell(i + 1, 0), _w1[0])
    put(_tb1.cell(i + 1, 0), '%.3f' % ifv)
    for j, val in enumerate(I_A[i]):
        cell_setup(_tb1.cell(i + 1, j + 1), _w1[j + 1])
        put(_tb1.cell(i + 1, j + 1), '%g' % val)
row_height(_tb1, 0, 620)

figure(GEN + r'\figures_tungsten\fig_w1.png', 14.30,
       '图 3  不同灯丝电流下的阳极电流—阳极电压曲线')

caption('表 2  不同温度 T 和 √U_a 下对应的阳极电流的对数值 lg I_a')
_tb2 = new_table(9, 9, _w1)
cell_setup(_tb2.cell(0, 0), _w1[0])
put_lines(_tb2.cell(0, 0), ['√U_a', 'lg I_a', 'T / 10³K'], size=9)
for j, su in enumerate(SQRT_U):
    cell_setup(_tb2.cell(0, j + 1), _w1[j + 1])
    put(_tb2.cell(0, j + 1), '%.1f' % su, bold=True)
for i, tv in enumerate(T2['T_1000K']):
    cell_setup(_tb2.cell(i + 1, 0), _w1[0])
    put(_tb2.cell(i + 1, 0), '%.2f' % tv)
    for j, val in enumerate(LG_IA[i]):
        cell_setup(_tb2.cell(i + 1, j + 1), _w1[j + 1])
        put(_tb2.cell(i + 1, j + 1), '%.2f' % val)
cell_setup(_tb2.cell(8, 0), _w1[0])
put(_tb2.cell(8, 0), '2.16')
for j, val in enumerate(ROW_T216):
    cell_setup(_tb2.cell(8, j + 1), _w1[j + 1])
    put(_tb2.cell(8, j + 1), '%.2f' % val)
row_height(_tb2, 0, 620)
body('注：表 2 的印刷表格只有 7 行（T = 1.88～2.12×10³ K），T = 2.16×10³ K 一行'
     '为原始记录页在表格下方补记的数据，与表 1 中 I_f = 0.775 A 一行的对数值一一对应，'
     '故一并列入。', indent=False, size=10.5)

figure(GEN + r'\figures_tungsten\fig_w2.png', 14.30,
       '图 4  lg I_a–√U_a 关系及外推求零场电流的截距')

caption('表 3  lg(I_s/T²)～1/T 数据表')
_w3 = [1700] + [825] * 8
_tb3 = new_table(4, 9, _w3)
_rows = [('T / 10³K', T_LIST, '%.2f'),
         ('1/T（10⁻³K⁻¹）', INV_T, '%.2f'),
         ('lg I_s', LG_I, '%.2f'),
         ('lg(I_s/T²)', LG_IT2, '%.2f')]
for i, (lab, vals, fmt) in enumerate(_rows):
    cell_setup(_tb3.cell(i, 0), _w3[0])
    put(_tb3.cell(i, 0), lab, bold=(i == 0))
    for j, v in enumerate(vals):
        cell_setup(_tb3.cell(i, j + 1), _w3[j + 1])
        put(_tb3.cell(i, j + 1), fmt % v)

figure(GEN + r'\figures_tungsten\fig_w3.png', 14.30,
       '图 5  lg(I_s/T²)–1/T 关系及由斜率求逸出功')

body('数据处理的具体过程如下：', indent=False, bold=True, keep_next=True)
body('（1）把表 1 中的阳极电压开平方得到 √U_a，把阳极电流取常用对数得到 lg I_a，'
     '整理成表 2；再将表 1 的 8 个灯丝电流按 T = 920.0 + 1600 I_f 换算成对应的阴极温度。')
def num(v, fmt):
    return (fmt % v).replace('-', '−')


body('（2）对表 2 中每一个温度，以 √U_a 为横坐标、lg I_a 为纵坐标作最小二乘直线，'
     '并把它外推到 √U_a = 0，直线与纵轴的交点即该温度下的 lg I_s，结果见图 4。'
     '8 个温度对应的截距依次为 '
     + '、'.join('%.3f' % v for v in LG_IS) + '，'
     '换算回指数即得到零场电流 I_s（表 3 第 3 行按原始记录只保留了两位小数）。')
body('（3）以 1/T 为横坐标（单位 10⁻³ K⁻¹）、lg(I_s/T²) 为纵坐标作直线，见图 5，'
     '最小二乘拟合得斜率 m = %s、截距 b = %s、相关系数 r = %s（R² = %s）。'
     % (num(M_PLOT, '%.2f'), num(B_PLOT, '%.2f'), num(R_VAL, '%.4f'), '%.4f' % R2_VAL))
body('（4）斜率换算为以 K⁻¹ 为单位：m_真 = m×10³ = %s×10⁴ K⁻¹。'
     % num(M_TRUE / 1e4, '%.3f'))
body('（5）由式（3）的斜率关系 eφ = 2.303 k |m_真| 得：eφ = 2.303×%.4f×10⁻²³×%.3f×10⁴ '
     '= %.2f×10⁻¹⁹ J，即 φ = %.3f V，eφ = %.2f eV。'
     % (K_B * 1e23, abs(M_TRUE) / 1e4, PHI_J * 1e19, PHI_EV, PHI_EV))
body('（6）与公认值 eφ = 4.54 eV 相比，相对误差 E = (%.3f − 4.54)/4.54 = %s%%；'
     '由斜率的标准偏差 u(m) = %.2f 得 u(eφ) = %.2f eV，故测量结果可写作 '
     'eφ = (%.2f ± %.2f) eV，与公认值之差约等于 1.4 倍不确定度。'
     % (PHI_EV, num(REL_ERR, '%.1f'), SE_M, U_PHI_EV, PHI_EV, U_PHI_EV))
body('（7）需要说明的是，表 3 第 4 行是直接用 lg I_s 减去 2lg(T/10³) 得到的，'
     '比真实的 lg(I_s/T²) 系统性高出 6.000（= 2lg 1000）。这只是一个整体平移，'
     '最小二乘直线的斜率并不改变，因而对 eφ 的结果没有影响；若把温度直接以 K 代入重算，'
     '得到的斜率与上述完全一致。')

# ---------------------------------------------------------------- 体会
heading('分析及体会', '20%', NOTE_EXP)
para_seq(SECTIONS['体会'])

# ---------------------------------------------------------------- 思考题
for ln in SECTIONS['思考题']:
    if ln.startswith('解答'):
        body(ln, indent=False)
    elif ln.startswith('题目'):
        body(ln, indent=False, bold=True, keep_next=True)

# ---------------------------------------------------------------- 超星截图
heading('超星平台本实验在线练习得分截图（请竖排粘贴）', '5%')
new_p(keep_next=True)
frame = doc.add_table(rows=1, cols=1)
frame.alignment = WD_TABLE_ALIGNMENT.CENTER
frame.autofit = False
frame.columns[0].width = tw(8300)
frame.cell(0, 0).width = tw(8300)
_tcPr = frame.cell(0, 0)._tc.get_or_add_tcPr()
_b = OxmlElement('w:tcBorders')
for _edge in ('top', 'left', 'bottom', 'right'):
    _e = OxmlElement('w:' + _edge)
    _e.set(qn('w:val'), 'nil')
    _b.append(_e)
_tcPr.append(_b)
_cell = frame.cell(0, 0)
_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
_p0 = _cell.paragraphs[0]
_p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
_pf = _p0.paragraph_format
_pf.first_line_indent = Pt(0)
_pf.space_before = Pt(6)
_pf.space_after = Pt(6)
_p0.add_run().add_picture(SHOT, width=Cm(8.20))

for _t in doc.tables:
    kill_vmerge_inner_lines(_t)

for _sec in doc.sections:
    for _p in _sec.footer.paragraphs:
        for _r in _p.runs:
            if '长度测量' in _r.text:
                _r.text = _r.text.replace('长度测量', '钨的逸出功')
doc.save(OUT)

# ---------------------------------------------------------------- verify
lines = []
lines.append('file: %s' % OUT)
lines.append('paragraphs=%d tables=%d images=%d' %
             (len(doc.paragraphs), len(doc.tables), len(doc.inline_shapes)))
lines.append('header=%r footer=%r' % (
    doc.sections[0].header.paragraphs[0].text,
    doc.sections[0].footer.paragraphs[0].text))
lines.append('phi_eV=%.4f  phi_J=%.4e  rel_err=%.3f%%  u(eV)=%.4f' %
             (PHI_EV, PHI_J, REL_ERR, U_PHI_EV))
for sh in doc.inline_shapes:
    lines.append('  image %.2f x %.2f cm' % (sh.width.cm, sh.height.cm))
for t in doc.tables:
    w = sum((c.width.twips if c.width else 0) for c in t.rows[0].cells)
    lines.append('  table %dx%d width_sum=%s' % (len(t.rows), len(t.columns), w))
lines.append('--- paragraphs ---')
for i, p in enumerate(doc.paragraphs):
    sz = ''
    for r in p.runs:
        if r.font.size:
            sz = r.font.size.pt
            break
    lines.append('[%d] <%s> ind=%s | %s' % (
        i, sz, p.paragraph_format.first_line_indent,
        p.text[:110].replace('\n', ' ')))
lines.append('--- tables ---')
for ti, t in enumerate(doc.tables):
    lines.append('TABLE %d %dx%d' % (ti, len(t.rows), len(t.columns)))
    for r in t.rows:
        lines.append('  | ' + ' | '.join(c.text.replace('\n', '/') for c in r.cells))
open(GEN + r'\verify_tungsten.txt', 'w', encoding='utf-8').write('\n'.join(lines))
print('verify written', len('\n'.join(lines)))
