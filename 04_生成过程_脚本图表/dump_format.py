# -*- coding: utf-8 -*-
"""Dump detailed formatting of reference DOCX files to UTF-8 text."""
import sys, io
from docx import Document
from docx.oxml.ns import qn

OUT = r'C:\Users\lidey\Desktop\实验报告制作沙盒\format_dump.txt'
buf = io.StringIO()

def p(x):
    buf.write(str(x) + '\n')

def run_info(r):
    nm = r.font.name
    sz = r.font.size.pt if r.font.size else None
    ea = None
    if r._element.rPr is not None and r._element.rPr.rFonts is not None:
        ea = r._element.rPr.rFonts.get(qn('w:eastAsia'))
    return f"[{nm}|ea={ea}|{sz}pt|b={r.bold}|i={r.italic}|u={r.underline}]"

def dump(path, label):
    p('=' * 100)
    p(f'FILE {label}: {path}')
    d = Document(path)
    s = d.sections[0]
    p(f'page {s.page_width.cm:.2f}x{s.page_height.cm:.2f}cm  margins L{s.left_margin.cm:.2f} R{s.right_margin.cm:.2f} T{s.top_margin.cm:.2f} B{s.bottom_margin.cm:.2f}  header{s.header_distance.cm:.2f} footer{s.footer_distance.cm:.2f}')
    try:
        p(f'header text: {[pp.text for pp in s.header.paragraphs]}')
        p(f'footer text: {[pp.text for pp in s.footer.paragraphs]}')
    except Exception as e:
        p(f'header/footer err {e}')
    # style defaults
    st = d.styles['Normal']
    p(f'Normal style: font={st.font.name} size={st.font.size.pt if st.font.size else None} '
      f'spacing/before={st.paragraph_format.space_before} after={st.paragraph_format.space_after} '
      f'line={st.paragraph_format.line_spacing} rule={st.paragraph_format.line_spacing_rule}')
    dflt = d.styles.element.find(qn('w:docDefaults'))
    if dflt is not None:
        from lxml import etree
        p('docDefaults: ' + etree.tostring(dflt, encoding='unicode')[:1200])

    body = d.element.body
    tbl_i = 0
    par_i = 0
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    for child in body.iterchildren():
        if child.tag == qn('w:p'):
            par = Paragraph(child, d)
            pf = par.paragraph_format
            runs = ''.join(run_info(r) for r in par.runs[:3])
            nimg = len(child.findall('.//' + qn('w:drawing')))
            p(f'P{par_i} <{par.style.name}> al={par.alignment} ind1st={pf.first_line_indent} '
              f'indL={pf.left_indent} sb={pf.space_before} sa={pf.space_after} ls={pf.line_spacing}/{pf.line_spacing_rule} '
              f'runs={runs} nimg={nimg} :: {par.text}')
            par_i += 1
        elif child.tag == qn('w:tbl'):
            t = Table(child, d)
            p(f'--- TBL{tbl_i} style={t.style.name if t.style else None} rows={len(t.rows)} cols={len(t.columns)}')
            cf = t.rows[0].cells[0]
            p(f'    cell0 para fmt: {cf.paragraphs[0].alignment} ' + ''.join(run_info(r) for r in cf.paragraphs[0].runs[:2]))
            for ri, row in enumerate(t.rows):
                cells = []
                for c in row.cells:
                    cells.append(c.text.replace('\n', '\\n'))
                p(f'    R{ri}: ' + ' | '.join(cells))
            tbl_i += 1

for path, label in [
    (r'C:\Users\lidey\Desktop\实验报告制作沙盒\热膨胀系数.docx', 'A'),
    (r'C:\Users\lidey\Desktop\实验报告制作沙盒\霍尔效应-已填数据.docx', 'B'),
]:
    dump(path, label)

with open(OUT, 'w', encoding='utf-8') as f:
    f.write(buf.getvalue())
print('written', OUT, len(buf.getvalue()))
