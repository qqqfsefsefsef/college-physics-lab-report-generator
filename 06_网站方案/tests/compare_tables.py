# -*- coding: utf-8 -*-
r"""逐格对比「直接生成版」与「网页流水线版」两份 docx 的表格。"""
import io
import sys

from docx import Document

sys.stdout.reconfigure(encoding='utf-8')

A = r'C:\Users\lidey\Desktop\实验报告制作沙盒\05_成品_报告\光电效应的研究.docx'
B = r'C:\Users\lidey\Desktop\实验报告制作沙盒\07_流水线\projects\photo-demo\out\光电效应的研究.docx'


def grid(path):
    d = Document(path)
    out = []
    for t in d.tables:
        rows = []
        for r in t.rows:
            rows.append([c.text.strip().replace('\n', '|') for c in r.cells])
        out.append(rows)
    return d, out


da, ga = grid(A)
db, gb = grid(B)
w = io.StringIO()
w.write('A 段落=%d 表=%d\nB 段落=%d 表=%d\n\n' % (
    len(da.paragraphs), len(ga), len(db.paragraphs), len(gb)))

for i in range(max(len(ga), len(gb))):
    ta = ga[i] if i < len(ga) else None
    tb = gb[i] if i < len(gb) else None
    ra = '%dx%d' % (len(ta), len(ta[0])) if ta else '-'
    rb = '%dx%d' % (len(tb), len(tb[0])) if tb else '-'
    w.write('== 表%d  A=%s  B=%s %s\n' % (i + 1, ra, rb, 'OK' if ra == rb else '形状不同'))
    if ta and tb and len(ta) == len(tb) and len(ta[0]) == len(tb[0]):
        bad = 0
        for r in range(len(ta)):
            for c in range(len(ta[0])):
                # 新版把 U_s 排成真下标，cell.text 里下划线会消失 → 比较前统一去掉
                if ta[r][c].replace('_', '') != tb[r][c].replace('_', ''):
                    bad += 1
                    if bad <= 6:
                        w.write('   (r%d,c%d) A=%r B=%r\n' % (r, c, ta[r][c], tb[r][c]))
        w.write('   逐格不同 %d / %d\n' % (bad, len(ta) * len(ta[0])))
    if ta:
        w.write('   A 首行: %s\n' % (' | '.join(ta[0])))
    if tb:
        w.write('   B 首行: %s\n' % (' | '.join(tb[0])))
    w.write('\n')

open(r'C:\Users\lidey\Desktop\实验报告制作沙盒\06_网站方案\tests\diff_tables.txt',
     'w', encoding='utf-8').write(w.getvalue())
print('written')
