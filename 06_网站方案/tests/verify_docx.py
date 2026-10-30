import sys

from docx import Document

p = sys.argv[1]
d = Document(p)
txt = []
for par in d.paragraphs:
    t = par.text
    if t.strip():
        txt.append(t)
bad = [t for t in txt if ('[[' in t or '===' in t or '*' in t or 'None' in t)]
print('paragraphs=%d tables=%d chars=%d' % (len(d.paragraphs), len(d.tables), sum(len(t) for t in txt)))
print('bad=', bad[:5])
for i, t in enumerate(txt[:14]):
    print('  [%d] %s' % (i, t[:70]))
print('  ...')
for i, t in enumerate(txt[-10:]):
    print('  [end%d] %s' % (i, t[:70]))
for k, tb in enumerate(d.tables):
    print('table%d %dx%d  r0=%s' % (k, len(tb.rows), len(tb.columns),
                                     [c.text for c in tb.rows[0].cells][:5]))
