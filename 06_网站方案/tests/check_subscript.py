# -*- coding: utf-8 -*-
r"""找出 docx 里没被解析成上下标的标记（_ ^ { }）以及真实下标 run 的数量。"""
import io
import sys

from docx import Document

sys.stdout.reconfigure(encoding='utf-8')
P = r'C:\Users\lidey\Desktop\实验报告制作沙盒\05_成品_报告\光电效应的研究_网页流水线版.docx'
d = Document(P)

bad = []
sub = 0
sup = 0


def scan_run(txt, where, run):
    global sub, sup
    if run.font.subscript:
        sub += 1
    if run.font.superscript:
        sup += 1
    if any(ch in txt for ch in '_^{}'):
        bad.append((where, txt))


for i, p in enumerate(d.paragraphs):
    for r in p.runs:
        scan_run(r.text, 'P%d' % i, r)

for ti, t in enumerate(d.tables):
    for ri, row in enumerate(t.rows):
        for ci, c in enumerate(row.cells):
            for p in c.paragraphs:
                for r in p.runs:
                    scan_run(r.text, 'T%d(%d,%d)' % (ti + 1, ri, ci), r)

w = io.StringIO()
w.write('真下标 run: %d   真上标 run: %d\n' % (sub, sup))
w.write('仍带标记的 run: %d\n' % len(bad))
seen = set()
for where, txt in bad:
    if txt in seen:
        continue
    seen.add(txt)
    w.write('  %-12s %r\n' % (where, txt))
io.open(r'C:\Users\lidey\Desktop\实验报告制作沙盒\06_网站方案\tests\sub_check.txt',
        'w', encoding='utf-8').write(w.getvalue())
print('ok', len(bad))
