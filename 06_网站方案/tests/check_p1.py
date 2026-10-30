# -*- coding: utf-8 -*-
"""P1 验证：/style（读作业抽风格卡）+ /render 带学习通截图真贴进 docx。

用法：python check_p1.py [port]
"""
import base64
import json
import os
import subprocess
import sys
import time
import urllib.error
import urllib.request

HERE = os.path.dirname(os.path.abspath(__file__))
WORK = os.path.abspath(os.path.join(HERE, '..', '..'))
PIPE = os.path.join(WORK, '07_流水线')
OUT = os.path.join(WORK, '05_成品_报告')
SHOT = os.path.join(WORK, '01_素材_数据与截图',
                    'Screenshot_2026-10-26-22-18-26-754_com.chaoxing.m.jpg')
PORT = int(sys.argv[1]) if len(sys.argv) > 1 else 8734
BASE = 'http://127.0.0.1:%d' % PORT
PY = sys.executable

fails = []


def ok(name, cond, extra=''):
    print(('[OK] ' if cond else '[FAIL] ') + name + (('  ' + str(extra)) if extra else ''))
    if not cond:
        fails.append(name)


def get(path, timeout=30):
    with urllib.request.urlopen(BASE + path, timeout=timeout) as r:
        return r.status, r.read()


def post(path, obj, timeout=180):
    req = urllib.request.Request(BASE + path, data=json.dumps(obj).encode('utf-8'),
                                 headers={'Content-Type': 'application/json'})
    try:
        with urllib.request.urlopen(req, timeout=timeout) as r:
            return r.status, json.loads(r.read().decode('utf-8'))
    except urllib.error.HTTPError as e:
        return e.code, json.loads(e.read().decode('utf-8'))


def wait(n=80):
    for _ in range(n):
        try:
            if get('/health', timeout=2)[0] == 200:
                return
        except Exception:
            pass
        time.sleep(0.3)
    raise SystemExit('桥没起来')


def main():
    env = dict(os.environ)
    env['PYTHONIOENCODING'] = 'utf-8'
    p = subprocess.Popen([PY, os.path.join(PIPE, 'bridge.py'), '--port', str(PORT)],
                         cwd=PIPE, env=env, stdout=subprocess.PIPE, stderr=subprocess.STDOUT)
    try:
        wait()
        # ---- 1. 风格卡
        hw = os.path.join(OUT, '光电效应的研究.docx')
        b64 = base64.b64encode(open(hw, 'rb').read()).decode('ascii')
        st, card = post('/style', {'docx_b64': b64})
        ok('/style 200', st == 200, st)
        ok('返回了摘要', bool(card.get('summary')), card.get('summary', '')[:60])
        s = card.get('stats') or {}
        ok('抽到句长/段落/人称', all(k in s for k in ('avg_sentence', 'paras', 'first_person_para_ratio')),
           '%s 字 句长 %.1f 段 %d' % (s.get('chars'), s.get('avg_sentence') or -1, s.get('paras') or -1))
        ok('风格提示里有硬约束', any('12 个字' in h for h in card.get('prompt_hints') or []))
        st2, bad = post('/style', {'docx_b64': base64.b64encode(b'not a docx').decode('ascii')})
        ok('垃圾文件被拦下', st2 == 400 and '不像' in bad.get('error', ''), bad.get('error'))

        # ---- 2. 拿桩表
        b64img = base64.b64encode(open(os.path.join(HERE, 'fixture.jpg'), 'rb').read()).decode('ascii')
        st, rec = post('/recognize', {'exp_id': 'photo', 'mode': 'mock',
                                      'image_b64': b64img, 'mime': 'image/jpeg'})
        ok('拿到 photo 桩表', st == 200 and 't1' in rec.get('tables', {}), st)

        # ---- 3. 不带截图 → 6 图 + 空框
        st, r1 = post('/render', {'exp_id': 'photo', 'tables': rec['tables']})
        ok('render(无截图) 200', st == 200, r1.get('error'))
        f1 = os.path.join(PIPE, 'projects', r1['project'], 'out', r1['name'])
        from docx import Document
        d1 = Document(f1)
        n1 = len(d1.inline_shapes)
        ok('无截图时 6 张图', n1 == 6, n1)
        tail1 = d1.tables[-1]
        ok('末表是 1×1 粘贴框', len(tail1.rows) == 1 and len(tail1.columns) == 1)
        ok('空框里有提示文字', '粘贴' in tail1.cell(0, 0).text, tail1.cell(0, 0).text[:20])

        # ---- 4. 带截图 → 7 图，最后一张在框里
        shot_b64 = base64.b64encode(open(SHOT, 'rb').read()).decode('ascii')
        st, r2 = post('/render', {'exp_id': 'photo', 'tables': rec['tables'],
                                  'shot_b64': shot_b64, 'shot_mime': 'image/jpeg'})
        ok('render(带截图) 200', st == 200, r2.get('error'))
        f2 = os.path.join(PIPE, 'projects', r2['project'], 'out', r2['name'])
        d2 = Document(f2)
        n2 = len(d2.inline_shapes)
        ok('带截图时 7 张图', n2 == 7, n2)
        cell = d2.tables[-1].cell(0, 0)
        ok('粘贴框里真有图片', len(cell.paragraphs[0].runs) == 1 and
           cell.paragraphs[0].runs[0]._element.findall(
               '{http://schemas.openxmlformats.org/drawingml/2006/main}graphic') != [] or
           'graphic' in cell.paragraphs[0]._p.xml, cell.text[:10])
        ok('截图落盘在项目里', os.path.exists(os.path.join(
            PIPE, 'projects', r2['project'], 'source', 'shot.jpg')))
        ok('文件比无截图时大', os.path.getsize(f2) > os.path.getsize(f1),
           '%d B vs %d B' % (os.path.getsize(f2), os.path.getsize(f1)))
    finally:
        p.terminate()
        try:
            p.wait(timeout=10)
        except Exception:
            p.kill()
    print('\n[SUM] ' + ('全部通过' if not fails else '%d 项失败: %s' % (len(fails), '; '.join(fails))))
    return 0 if not fails else 1


if __name__ == '__main__':
    sys.exit(main())
